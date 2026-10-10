import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    BG_COLOR = RGBColor(251, 247, 238)        # #fbf7ee (kraft paper)
    CARD_BG = RGBColor(255, 255, 255)         # #ffffff
    PRIMARY = RGBColor(45, 55, 72)            # #2d3748 (dark slate)
    ACCENT_BLUE = RGBColor(49, 130, 206)      # #3182ce
    ACCENT_GOLD = RGBColor(214, 158, 46)      # #d69e2e
    ACCENT_GREEN = RGBColor(47, 133, 90)      # #2f855a
    TEXT_MUTED = RGBColor(113, 128, 150)      # #718096
    BORDER_COLOR = RGBColor(45, 55, 72)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

    def add_header(slide, title_text, category_text="Хакатон «ИТ-Прорыв» | Номинация «Мобильная разработка»"):
        # Category pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

    def add_card(slide, left, top, width, height, title="", border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        
        if title:
            tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(width - 0.5), Inches(0.5))
            p = tb.text_frame.paragraphs[0]
            p.text = title
            p.font.size = Pt(16)
            p.font.bold = True
            p.font.color.rgb = PRIMARY
        return card

    # =========================================================================
    # СЛАЙД 0: ТИТУЛЬНЫЙ
    # =========================================================================
    s0 = prs.slides.add_slide(blank_layout)
    set_slide_background(s0)

    # Big Card in center
    add_card(s0, 1.2, 1.0, 10.933, 5.5, border_color=PRIMARY)

    tb = s0.shapes.add_textbox(Inches(1.8), Inches(1.5), Inches(9.7), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "ВСЕРОССИЙСКИЙ ХАКАТОН «ИТ-ПРОРЫВ»"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p2 = tf.add_paragraph()
    p2.text = "CyberStack: Инженер Памяти"
    p2.font.size = Pt(36)
    p2.font.bold = True
    p2.font.color.rgb = PRIMARY

    p3 = tf.add_paragraph()
    p3.text = "Игровое интерактивное PWA-приложение для изучения структуры данных «Стек» и принципа LIFO для подростков 12–17 лет"
    p3.font.size = Pt(17)
    p3.font.color.rgb = TEXT_MUTED

    p4 = tf.add_paragraph()
    p4.text = "\nНоминация: «Мобильная разработка» | Кейс: «Мобильное приложение для изучения программирования»\nРазработчик: Школьная команда «CyberStack»"
    p4.font.size = Pt(13)
    p4.font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 1: НАЗВАНИЕ И ИДЕЯ ПРИЛОЖЕНИЯ
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)
    add_header(s1, "1. Название и идея приложения")

    add_card(s1, 0.8, 1.8, 5.6, 5.0, "Проблема обучения")
    tb = s1.shapes.add_textbox(Inches(1.05), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Перегруз скучной теорией:\nНачинающие программисты быстро теряют интерес из-за обилия сухого текста и отсутствия наглядности."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY
    p = tf.add_paragraph()
    p.text = "\n• Однотипные задания:\nБольшинство приложений сводятся к банальным тестам с выбором 1 варианта из 4."
    p.font.size = Pt(14)
    p.font.color.rgb = PRIMARY
    p2 = tf.add_paragraph()
    p2.text = "\n• Отложенный результат:\nПодросток не понимает, зачем нужна изучаемая конструкция в реальной разработке."
    p2.font.size = Pt(14)
    p2.font.color.rgb = PRIMARY

    add_card(s1, 6.9, 1.8, 5.6, 5.0, "Наше решение: CyberStack", border_color=ACCENT_BLUE)
    tb = s1.shapes.add_textbox(Inches(7.15), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Формат «Лабораторного скетчбука»:\nИнтерфейс на легком фреймворке PaperCSS воспринимается как интерактивная тетрадь юного изобретателя."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY
    p = tf.add_paragraph()
    p.text = "\n• Физический интерактивный симулятор:\nУченик своими руками пушит (PUSH) и достает (POP) блоки памяти в анимированном реакторе стека."
    p.font.size = Pt(14)
    p.font.color.rgb = PRIMARY
    p2 = tf.add_paragraph()
    p2.text = "\n• Робот-напарник Байт:\nБортовой ИИ-ассистент, готовый объяснить термины простыми словами и дать контекстную подсказку."
    p2.font.size = Pt(14)
    p2.font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 2: ЦЕЛЕВАЯ АУДИТОРИЯ (12–17 ЛЕТ) И ПОТРЕБНОСТИ
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "2. Выбранная возрастная группа и её потребности")

    add_card(s2, 0.8, 1.8, 3.6, 5.0, "Портрет аудитории")
    tb = s2.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(3.2), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "Подростки 12–17 лет:\n• Школьники 7–11 классов\n• Изучают информатику и готовятся к ОГЭ/ЕГЭ\n• Интересуются геймдевом и алгоритмами\n• Привыкли к быстрым интерактивным приложениям"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    add_card(s2, 4.8, 1.8, 3.6, 5.0, "Ключевые потребности")
    tb = s2.shapes.add_textbox(Inches(5.0), Inches(2.5), Inches(3.2), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "Что важно подростку:\n• Короткие уроки (3–5 мин)\n• Понятный язык без академического занудства\n• Игровой отклик на действия\n• Безопасность: без регистрации и ввода личных данных"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    add_card(s2, 8.8, 1.8, 3.7, 5.0, "Ответ в CyberStack", border_color=ACCENT_GREEN)
    tb = s2.shapes.add_textbox(Inches(9.0), Inches(2.5), Inches(3.3), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "Реализация в проекте:\n• Модульная структура микрообучения\n• Реальные примеры (игры, Ctrl+Z, компилятор)\n• Награды (XP, бейджи, ранги)\n• PWA: мгновенный старт без учетных записей"
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 3: ТЕМА ПРОГРАММИРОВАНИЯ
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "3. Тема: Структура данных Стек и принцип LIFO")

    add_card(s3, 0.8, 1.8, 5.6, 5.0, "Почему это небанально?")
    tb = s3.shapes.add_textbox(Inches(1.05), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Отказ от банальных тем:\nБольшинство проектов берут переменные или циклы. Стек — это фундамент алгоритмов и олимпиадной информатики.\n\n• Профильный школьный уровень:\nТема входит в программу углубленной информатики и заданий ЕГЭ (анализ вложенных конструкций, вызовы функций).\n\n• Визуальная наглядность:\nСтек идеально визуализируется физической моделью стопки."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    add_card(s3, 6.9, 1.8, 5.6, 5.0, "Практическое применение в IT", border_color=ACCENT_BLUE)
    tb = s3.shapes.add_textbox(Inches(7.15), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "Где работает стек в реальном мире:\n\n1. История действий (Undo / Ctrl+Z):\nОтмена изменений в Figma, Фотошопе и текстовых процессорах.\n\n2. Навигация браузера:\nКнопки «Назад» и «Вперед» сохраняют адреса в стеке истории.\n\n3. Анализ синтаксиса компилятором:\nПроверка баланса скобок {} [] () в средах разработки (VS Code, PyCharm)."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 4: СТРУКТУРА ТРЁХ УРОКОВ
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "4. Структура образовательного маршрута")

    lessons = [
        ("Урок 1: Башня Памяти", "Основы LIFO, операции PUSH и POP", "• Аналогия со стопкой тарелок\n• Команды PUSH (положить) и POP (снять)\n• Интерактивная песочница LIFO\n• Практика: сборка чипов для робота", 0.8),
        ("Урок 2: Машина Времени", "Стек отмены (Ctrl+Z) в редакторах", "• Как программы запоминают действия\n• Почему отмена требует реверса (LIFO)\n• Практика: восстановление цепочки отмены при сбое манипулятора", 4.8),
        ("Урок 3: Детектив Скобок", "Как компилятор ищет SyntaxError", "• Алгоритм проверки парности скобок\n• Сопоставление открывающих и закрывающих\n• Практика: нахождение синтаксического бага в квантовом ядре", 8.8)
    ]

    for title, sub, desc, left in lessons:
        add_card(s4, left, 1.8, 3.7, 5.0, title)
        tb = s4.shapes.add_textbox(Inches(left + 0.2), Inches(2.4), Inches(3.3), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = sub
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = ACCENT_BLUE
        p = tf.add_paragraph()
        p.text = "\n" + desc
        p.font.size = Pt(13)
        p.font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 5: ФОРМАТЫ ПРАКТИЧЕСКИХ ЗАДАНИЙ
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "5. Разнообразие практических заданий")

    formats = [
        ("Формат 1: Интерактивный сборщик", "Урок 1", "Ученик управляет пультом команд (+PUSH Синий, +PUSH Золотой, -POP). Задача — загрузить в стек чипы памяти в строгом порядке снизу вверх, учитывая правило LIFO.", 0.8),
        ("Формат 2: Восстановление порядка", "Урок 2", "Манипулятор выполнил 4 действия и застрял. Ученик кликает по карточкам действий в порядке их отмены, восстанавливая безопасную последовательность Undo.", 4.8),
        ("Формат 3: Синтаксический детектив", "Урок 3", "Дана строка кода с ошибкой валидации скобок {[()]}. Ученик анализирует вершину стека и определяет, почему компилятор забраковал символ и какой тип ожидался.", 8.8)
    ]

    for title, sub, desc, left in formats:
        add_card(s5, left, 1.8, 3.7, 5.0, title)
        tb = s5.shapes.add_textbox(Inches(left + 0.2), Inches(2.4), Inches(3.3), Inches(4.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = sub
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.bold = True
        tf.paragraphs[0].font.color.rgb = ACCENT_GREEN
        p = tf.add_paragraph()
        p.text = "\n" + desc
        p.font.size = Pt(13)
        p.font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 6: ГЕЙМИФИКАЦИЯ И СИСТЕМА ПРОГРЕССА
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "6. Игровая механика и система прогресса")

    add_card(s6, 0.8, 1.8, 5.6, 5.0, "Игровые стимулы")
    tb = s6.shapes.add_textbox(Inches(1.05), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Очки опыта (XP):\nНачисляются за успешное решение задач (+60, +70, +80 XP).\n\n• Инженерные ранги:\nСтажёр Блокнота → Оператор Стека → Мастер LIFO.\n\n• Коллекция из 6 бейджей:\n«Первый PUSH», «Чистый Код», «Мастер Ctrl+Z», «Парсер Скобок», «Диплом LIFO», «Друг Байта».\n\n• Именной сертификат:\nГенерируется на финальном экране со штампом ОТК."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    add_card(s6, 6.9, 1.8, 5.6, 5.0, "Мультимедиа и интерактив", border_color=ACCENT_GOLD)
    tb = s6.shapes.add_textbox(Inches(7.15), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• Физический симулятор стека:\nАнимированное падение чипов в стакан памяти и вылет при POP.\n\n• Процедурный звук (Web Audio API):\nСинтезирует приятные клики, всплытие чипа, победный аккорд и фанфары прямо в коде (без сторонних mp3 файлов).\n\n• Управление звуком:\nКнопка быстрого включения/отключения аудиоэффектов."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 7: ПОЛЬЗОВАТЕЛЬСКИЙ ПУТЬ (USER JOURNEY)
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "7. Пользовательский путь (User Journey)")

    stages = [
        ("1. Старт", "Стартовый экран с идеей и кнопкой начала обучения."),
        ("2. Онбординг", "Вводный гайд: правила, механики наград и напарник Байт."),
        ("3. Карта", "Индикация прогресса (0-100%), выбор доступного урока."),
        ("4. Теория", "Краткое объяснение + песочница LIFO для закрепления."),
        ("5. Практика", "Выполнение задания, проверка, подсказки при ошибке."),
        ("6. Финал", "Сертификат инженера, итоговые XP и бейджи.")
    ]

    for i, (title, desc) in enumerate(stages):
        row = i // 3
        col = i % 3
        left = 0.8 + col * 4.0
        top = 1.8 + row * 2.5
        add_card(s7, left, top, 3.7, 2.2, title)
        tb = s7.shapes.add_textbox(Inches(left + 0.2), Inches(top + 0.6), Inches(3.3), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 8: ТЕХНОЛОГИИ И ИИ-ПОМОЩНИК
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "8. Использованные технологии и ИИ-помощник")

    add_card(s8, 0.8, 1.8, 5.6, 5.0, "Технологический стек")
    tb = s8.shapes.add_textbox(Inches(1.05), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• PWA (Progressive Web App):\nУстанавливается на экран смартфона, работает офлайн благодаря Service Worker (кэш Cache-First).\n\n• PaperCSS + Vanilla CSS:\nСтиль лабораторного скетчбука, кастомные анимации падающих чипов, адаптивность.\n\n• Чистый JavaScript (ES6):\nПонятный модульный код (app.js, state.js, lessons.js, bot.js, audio.js).\n\n• Web Storage (localStorage):\nСохранение прогресса без серверов и баз данных."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    add_card(s8, 6.9, 1.8, 5.6, 5.0, "Бортовой робот Байт", border_color=ACCENT_BLUE)
    tb = s8.shapes.add_textbox(Inches(7.15), Inches(2.5), Inches(5.1), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = "• 100% Автономность:\nРаботает полностью офлайн на локальном семантическом движке правил и базы знаний.\n\n• Контекстные подсказки:\nЗнает, на каком уроке находится ученик, и помогает разобрать текущую задачу.\n\n• Защита персональных данных:\nНе запрашивает и не сохраняет личные данные (требование кейса).\n\n• Оффтоп-контроль:\nВежливо возвращает ученика к алгоритмам при посторонних вопросах."
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = PRIMARY

    # =========================================================================
    # СЛАЙД 9: СКРИНШОТЫ РАБОТАЮЩЕГО ПРОТОТИПА
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "9. Интерфейс работающего прототипа")

    screens = [
        ("Старт и Карта", "Стартовый экран и наглядная карта с индикатором прогресса уроков и блокнотом наград.", 0.8),
        ("Урок и Симулятор", "Интерактивный стакан памяти: кнопки PUSH и POP визуально наполняют стек чипами.", 4.8),
        ("ИИ-Бот и Финал", "Выдвижной диалог с роботом Байтом и итоговый экран с именным сертификатом инженера.", 8.8)
    ]

    for title, desc, left in screens:
        add_card(s9, left, 1.8, 3.7, 5.0, title)
        tb = s9.shapes.add_textbox(Inches(left + 0.2), Inches(2.5), Inches(3.3), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(14)
        tf.paragraphs[0].font.color.rgb = PRIMARY
        p = tf.add_paragraph()
        p.text = "\n• Адаптировано под мобильные экраны (360–430px)\n• Крупные тач-кнопки\n• Плавные анимации и звуковой отклик"
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_MUTED

    # =========================================================================
    # СЛАЙД 10: КОМАНДА И РАСПРЕДЕЛЕНИЕ РОЛЕЙ
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "10. Команда и распределение ролей")

    roles = [
        ("Капитан & Методист", "Проработка темы «Стек и LIFO», методическая адаптация теории для возраста 12–17 лет, тексты уроков и формулировка заданий."),
        ("Frontend-разработчик", "Верстка интерфейса на HTML5, адаптация PaperCSS, реализация PWA-манифеста и сервисного воркера офлайн-доступа."),
        ("JS & Game Logic", "Логика управления состоянием (state.js), физический симулятор стакана стека, проверка заданий и Web Audio API."),
        ("AI & UX / Тестирование", "Разработка базы знаний и логики робота Байта, сценарий видеоролика, мобильное тестирование и подготовка презентации.")
    ]

    for i, (role, desc) in enumerate(roles):
        row = i // 2
        col = i % 2
        left = 0.8 + col * 6.0
        top = 1.8 + row * 2.5
        add_card(s10, left, top, 5.7, 2.2, role, border_color=ACCENT_BLUE if i==0 else BORDER_COLOR)
        tb = s10.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.65), Inches(5.2), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.paragraphs[0].text = desc
        tf.paragraphs[0].font.size = Pt(13)
        tf.paragraphs[0].font.color.rgb = PRIMARY

    prs.save("presentation_cyberstack.pptx")
    print("Презентация успешно создана: presentation_cyberstack.pptx")

if __name__ == "__main__":
    create_deck()
